import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC = "data/powerbi"
OUT = f"{SRC}/scouting_dashboard.xlsx"

PROFILE_NAMES = {
    "player": "Player", "team": "Team", "league": "League", "age": "Age",
    "season": "Season", "n_seasons": "Seasons Played", "position_group": "Position Group",
    "minutes": "Minutes", "role_1_name": "Role", "role_2_name": "Second Role",
    "role_gap": "Role Gap", "tm_sub_position": "Position", "height_in_cm": "Height (cm)",
    "foot": "Foot", "value_m": "Value (EURm)", "value_ratio": "Value Ratio",
    "contract_expiration_date": "Contract Expiry", "contract_years": "Contract Years Left",
    "tm_current_club": "Current Club",
    "goals_per90": "Goals per 90", "assists_per90": "Assists per 90",
    "shots_per90": "Shots per 90", "crosses_per90": "Crosses per 90",
    "interceptions_per90": "Interceptions per 90", "tackles_won_per90": "Tackles Won per 90",
    "goals_pct": "Goals Percentile", "assists_pct": "Assists Percentile",
    "shots_pct": "Shots Percentile", "crosses_pct": "Crosses Percentile",
    "interceptions_pct": "Interceptions Percentile", "tackles_won_pct": "Tackles Won Percentile",
}
SIMILAR_NAMES = {
    "player": "Player", "rank": "Rank", "similar_player": "Similar Player",
    "similar_team": "Similar Team", "similar_league": "Similar League",
    "similar_age": "Similar Age", "similar_sub_position": "Similar Position",
    "similar_value_m": "Similar Value (EURm)", "similar_value_ratio": "Similar Value Ratio",
    "similar_contract_years": "Similar Contract Years Left", "similar_role": "Similar Role",
    "distance": "Distance",
}


def bands(df):
    """Ordered text bands (numbered so they sort correctly in Power BI)."""
    df["Age Band"] = pd.cut(df["Age"], [0, 21, 24, 27, 30, 60],
                            labels=["1: 21 and under", "2: 22-24", "3: 25-27", "4: 28-30", "5: 31+"]).astype(str)
    df["Value Band"] = pd.cut(df["Value (EURm)"], [-1, 2, 5, 15, 40, 1e4],
                              labels=["1: under 2M", "2: 2-5M", "3: 5-15M", "4: 15-40M", "5: 40M+"]).astype(str)
    df["Ratio Band"] = pd.cut(df["Value Ratio"], [0, 0.5, 0.8, 1.25, 2, 1e4],
                              labels=["1: well below model (0.5 or less)", "2: below (0.5-0.8)",
                                      "3: fair (0.8-1.25)", "4: above (1.25-2)",
                                      "5: well above (over 2)"]).astype(str)
    for c in ("Age Band", "Value Band", "Ratio Band"):
        df[c] = df[c].replace("nan", "Unknown")
    return df


def flags(row, single_season_flag=True):
    out = []
    if single_season_flag and row["Seasons Played"] == 1:
        out.append("one season only")
    if row["Minutes"] < 1500:
        out.append("low minutes")
    if row["Role Gap"] <= 0.3:
        out.append(f"hybrid ({row['Second Role']})")
    if row["Position"] in ("Central Midfield", "Defensive Midfield") and row["Crosses per 90"] >= 3.0:
        out.append("crosses may be set pieces")
    cy = row["Contract Years Left"]
    if pd.notna(cy):
        if cy < 0:
            out.append("contract expired per July 2026 data (verify)")
        elif cy <= 1:
            out.append("contract ends within 12 months")
    if pd.isna(row["Value (EURm)"]):
        out.append("no market value")
    return "; ".join(out)


def prep_profiles(path, single_season_flag):
    df = pd.read_csv(path).rename(columns=PROFILE_NAMES)
    df["Contract Expiry"] = pd.to_datetime(df["Contract Expiry"])
    df["Age"] = df["Age"].astype("Int64")
    df = bands(df)
    df["Flags"] = df.apply(flags, axis=1, single_season_flag=single_season_flag)
    return df


def write_table(writer, df, sheet, table_name):
    df.to_excel(writer, sheet_name=sheet, index=False)
    ws = writer.sheets[sheet]
    ref = f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    for i, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(col)) + 2, 12), 34)
    ws.freeze_panes = "B2"


def main():
    profiles = prep_profiles(f"{SRC}/player_profiles.csv", True)
    seasons = prep_profiles(f"{SRC}/player_season_profiles.csv", False)
    similar = pd.read_csv(f"{SRC}/similar_players.csv").rename(columns=SIMILAR_NAMES)
    similar["Similar Age"] = similar["Similar Age"].astype("Int64")

    with pd.ExcelWriter(OUT, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        write_table(writer, profiles, "Profiles", "Profiles")
        write_table(writer, seasons, "Seasons", "Seasons")
        write_table(writer, similar, "Similar", "Similar")
    print(f"Wrote {OUT}")
    print(f"  Profiles: {len(profiles)} rows, {profiles.shape[1]} columns")
    print(f"  Seasons:  {len(seasons)} rows")
    print(f"  Similar:  {len(similar)} rows")


if __name__ == "__main__":
    main()